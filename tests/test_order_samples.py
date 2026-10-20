from decimal import Decimal
from importlib import import_module
from io import BytesIO
from types import SimpleNamespace

from django.apps import apps
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db import connection
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from operations.forms import ProgramOrderForm
from operations.models import Client, Inventory, InventoryBucket, Machine, Movement, Part, ProductionClose, ProductionOrder, WorkInProcess
from operations.services import create_program_order, edit_production_order, start_production, close_production


class OrderSamplesTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser('samples-admin', password='test-password')
        self.customer = Client.objects.create(code='SAMPLES', name='Samples client', external_id='S01')
        self.part = Part.objects.create(number='SAMPLE-001', client=self.customer, unit_weight_kg=Decimal('0.25'))
        self.client.force_login(self.user)

    def create(self, **kwargs):
        return create_program_order(client_name=self.customer.name, client=self.customer,
            part_number=kwargs.pop('part_number', self.part.number), program='SAMPLE-PROGRAM',
            quantity=kwargs.pop('quantity', 20), user=self.user, **kwargs)

    def test_same_part_can_have_normal_and_sample_orders(self):
        normal = self.create()
        sample = self.create(is_sample=True, quantity=999)
        self.assertEqual(normal.part_id, sample.part_id)
        self.assertEqual(normal.quantity, 20)
        self.assertEqual(sample.quantity, 5)
        self.assertEqual(sample.remaining_quantity, 5)
        self.assertEqual(sample.part_display, 'SAMPLE-001 (MUESTRAS)')
        self.assertEqual(Movement.objects.get(folio=sample.folio).quantity, 5)
        self.assertEqual(Part.objects.count(), 1)

    def test_single_form_and_view_enforce_five_without_javascript(self):
        payload = {'client': self.customer.name, 'part_number': self.part.number,
                   'program': 'WEB-SAMPLE', 'is_sample': 'on'}
        form = ProgramOrderForm(payload)
        self.assertTrue(form.is_valid(), form.errors)
        self.assertEqual(form.cleaned_data['quantity'], 5)
        self.assertEqual(self.client.post(reverse('load-program'), payload).status_code, 302)
        self.assertEqual(ProductionOrder.objects.get().quantity, 5)

    def test_bulk_preview_sample_choice_is_server_enforced_and_idempotent(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        ws.append(['ID Cliente', 'Orden de Produccion', 'Linea Prod Clte', 'Fecha de Entrega', 'Num. Parte', 'Cantidad', 'Linea'])
        ws.append(['S01', 'BULK-SAMPLE', '', None, self.part.number, 50, ''])
        ws.append(['S01', 'BULK-NORMAL', '', None, self.part.number, 60, ''])
        stream = BytesIO()
        wb.save(stream)
        preview = self.client.post(reverse('bulk-load-program'), {'file': SimpleUploadedFile('samples.xlsx', stream.getvalue())})
        token = preview.context['preview_token']
        self.assertTrue(token)
        self.assertContains(preview, 'form="bulk-confirm-form"')
        payload = {'action': 'confirm', 'preview_token': token, 'sample_2': 'on', 'quantity_2': 12345}
        self.assertEqual(self.client.post(reverse('bulk-load-program'), payload).status_code, 302)
        self.client.post(reverse('bulk-load-program'), payload)
        self.assertEqual(ProductionOrder.objects.count(), 2)
        sample = ProductionOrder.objects.get(program='BULK-SAMPLE')
        normal = ProductionOrder.objects.get(program='BULK-NORMAL')
        self.assertTrue(sample.is_sample)
        self.assertEqual(sample.quantity, 5)
        self.assertFalse(normal.is_sample)
        self.assertEqual(normal.quantity, 60)

    def test_edit_classification_preserves_existing_quantities_and_production(self):
        order = self.create(quantity=16)
        machine = Machine.objects.create(code='SAMPLE-MACHINE')
        work = start_production(order=order, machine=machine, quantity=6, employee=self.user)
        close = close_production(work_item=work, quantity=3, employee=self.user)
        order.refresh_from_db()
        original_remaining = order.remaining_quantity
        edited = edit_production_order(order_id=order.pk, program=order.program, part=self.part,
            quantity=order.quantity, is_sample=True, user=self.user)
        self.assertEqual(edited.quantity, 16)
        self.assertEqual(edited.remaining_quantity, original_remaining)
        self.assertEqual(ProductionClose.objects.get(pk=close.pk).quantity, 3)
        self.assertEqual(Movement.objects.get(folio=order.folio).quantity, 16)
        response = self.client.get(reverse('order-list'), {'is_sample': 'true'})
        self.assertContains(response, 'SAMPLE-001 (MUESTRAS)')
        self.assertNotContains(self.client.get(reverse('order-list'), {'is_sample': 'false'}), order.folio)

    def test_legacy_linking_preserves_every_existing_field_and_unmatched_parts(self):
        alias = Part.objects.create(number='SAMPLE-001 (MUESTRAS)', client=self.customer, unit_weight_kg=Decimal('0.10'))
        different = Client.objects.create(code='DIFFERENT', name='Different')
        other_base = Part.objects.create(number='OTHER-001', client=different)
        conflict = Part.objects.create(number='OTHER-001 (MUESTRAS)', client=self.customer)
        unmatched = Part.objects.create(number='UNMATCHED (MUESTRAS)', client=self.customer)
        legacy = ProductionOrder.objects.create(folio='OLD-SAMPLE', program='OLD', part=alias, quantity=16, remaining_quantity=11)
        untouched = ProductionOrder.objects.create(folio='OLD-UNMATCHED', program='OLD', part=unmatched, quantity=9, remaining_quantity=9)
        inv = Inventory.objects.create(part=alias, real=7, surplus=2)
        bucket = InventoryBucket.objects.create(inventory=inv, kind='PROGRAM', name='OLD', quantity=5)
        before = ProductionOrder.objects.values().get(pk=legacy.pk)
        import_module('operations.migrations.0017_order_samples').link_legacy_samples(apps, SimpleNamespace(connection=connection))
        legacy.refresh_from_db()
        after = ProductionOrder.objects.values().get(pk=legacy.pk)
        self.assertTrue(after.pop('is_sample'))
        before.pop('is_sample')
        self.assertEqual(before, after)
        alias.refresh_from_db()
        self.assertEqual(alias.sample_base_id, self.part.pk)
        self.assertEqual(alias.unit_weight_kg, Decimal('0.10'))
        self.assertEqual(legacy.part_id, alias.pk)
        self.assertEqual(Inventory.objects.get(pk=inv.pk).real, 7)
        self.assertEqual(InventoryBucket.objects.get(pk=bucket.pk).quantity, 5)
        conflict.refresh_from_db()
        unmatched.refresh_from_db()
        untouched.refresh_from_db()
        self.assertIsNone(conflict.sample_base_id)
        self.assertIsNone(unmatched.sample_base_id)
        self.assertFalse(untouched.is_sample)
        new = self.create(part_number=alias.number)
        self.assertEqual(new.part_id, self.part.pk)
        self.assertTrue(new.is_sample)
        self.assertEqual(new.quantity, 5)

    def test_exports_use_order_sample_label(self):
        sample = self.create(is_sample=True)
        response = self.client.get(reverse('download-filtered-orders'))
        wb = load_workbook(BytesIO(response.content), read_only=True, data_only=True)
        self.assertEqual(wb.active.cell(2, 6).value, sample.part_display)

    def test_catalog_hides_legacy_aliases_and_counts_visible_parts(self):
        self.part.in_universe_ramos = True
        self.part.save(update_fields=['in_universe_ramos'])
        alias = Part.objects.create(number='SAMPLE-001 (MUESTRAS)', client=self.customer, sample_base=self.part)
        unmatched = Part.objects.create(number='UNMATCHED (MUESTRAS)', client=self.customer)
        order = ProductionOrder.objects.create(folio='KEPT-HISTORY', program='OLD', part=alias,
            is_sample=True, quantity=16, remaining_quantity=11)
        url = reverse('universe')
        response = self.client.get(url, {'catalog': 'all'})
        self.assertEqual(response.context['catalog_total'], 2)
        self.assertEqual(response.context['universe_total'], 2)
        self.assertEqual({part.pk for part in response.context['parts']}, {self.part.pk, unmatched.pk})
        self.assertNotContains(response, alias.number)
        self.assertContains(response, unmatched.number)
        self.assertContains(self.client.get(url), 'Ver todos los números (2)')
        self.assertNotContains(self.client.get(url, {'catalog': 'all', 'q': alias.number}), alias.number + '</b>')
        order.refresh_from_db()
        self.assertEqual(order.part_id, alias.pk)
        self.assertEqual(order.quantity, 16)
        self.assertEqual(order.remaining_quantity, 11)
        self.assertEqual(order.part_display, 'SAMPLE-001 (MUESTRAS)')
        self.assertTrue(Part.objects.filter(pk=alias.pk).exists())

    def test_promoting_new_sample_reuses_id_extracts_diameter_and_preserves_history(self):
        part = Part.objects.create(number='NEW-008 (MUESTRAS) 3/8', client=self.customer, unit_weight_kg=Decimal('0.125'))
        original_part = Part.objects.values().get(pk=part.pk)
        order = ProductionOrder.objects.create(folio='OLD-NEW-008', program='HISTORY', part=part, quantity=16, remaining_quantity=7)
        original_order = ProductionOrder.objects.values().get(pk=order.pk)
        inventory = Inventory.objects.create(part=part, real=9, surplus=4)
        import_module('operations.migrations.0018_normalize_new_sample_parts').normalize_new_samples(apps, SimpleNamespace(connection=connection))
        part.refresh_from_db()
        order.refresh_from_db()
        self.assertEqual(part.number, 'NEW-008')
        self.assertEqual(part.diameter, '3/8')
        self.assertEqual(part.legacy_number, original_part['number'])
        expected_part = {**original_part, 'number': 'NEW-008', 'diameter': '3/8', 'legacy_number': original_part['number']}
        self.assertEqual(Part.objects.values().get(pk=part.pk), expected_part)
        self.assertEqual(ProductionOrder.objects.values().get(pk=order.pk), {**original_order, 'is_sample': True})
        self.assertEqual(Inventory.objects.get(pk=inventory.pk).real, 9)
        self.assertEqual(order.part_display, 'NEW-008 (MUESTRAS)')
        self.assertEqual(Part.objects.filter(number='NEW-008').count(), 1)
        # Historic Excel references resolve to this same record instead of creating a duplicate.
        new = self.create(part_number=original_part['number'])
        self.assertEqual(new.part_id, part.pk)
        self.assertTrue(new.is_sample)
        self.assertEqual(new.quantity, 5)
        normal = self.create(part_number=part.number, is_sample=False, quantity=12)
        self.assertEqual(normal.part_id, part.pk)
        self.assertFalse(normal.is_sample)
        self.assertEqual(normal.quantity, 12)

    def test_new_sample_conversion_skips_conflicts_and_does_not_guess_diameter(self):
        conflict = Part.objects.create(number='SAMPLE-001 (MUESTRAS)', client=self.customer)
        unknown = Part.objects.create(number='BRAND-NEW (MUESTRAS)', client=self.customer)
        import_module('operations.migrations.0018_normalize_new_sample_parts').normalize_new_samples(apps, SimpleNamespace(connection=connection))
        conflict.refresh_from_db()
        unknown.refresh_from_db()
        self.assertEqual(conflict.number, 'SAMPLE-001 (MUESTRAS)')
        self.assertEqual(conflict.legacy_number, '')
        self.assertEqual(unknown.number, 'BRAND-NEW')
        self.assertEqual(unknown.diameter, '')

    def test_user_confirmed_wrong_client_samples_hide_without_reassigning_history(self):
        daikin = Client.objects.create(code='DAIKIN', name='DAIKIN')
        rheem = Client.objects.create(code='RHEEM', name='RHEEM CIP-3501')
        original = Part.objects.create(number='3P848961-1', client=daikin)
        duplicate = Part.objects.create(number='3P848961-1 (MUESTRAS)', client=rheem)
        order = ProductionOrder.objects.create(folio='WRONG-CLIENT-HISTORY', program='OLD', part=duplicate,
            quantity=16, remaining_quantity=11)
        before = ProductionOrder.objects.values().get(pk=order.pk)
        import_module('operations.migrations.0019_hide_incorrect_client_samples').hide_incorrect_client_samples(apps, SimpleNamespace(connection=connection))
        duplicate.refresh_from_db()
        order.refresh_from_db()
        self.assertEqual(duplicate.sample_base_id, original.pk)
        self.assertEqual(duplicate.client_id, rheem.pk)
        self.assertEqual(ProductionOrder.objects.values().get(pk=order.pk), {**before, 'is_sample': True})
        self.assertEqual(order.part_display, '3P848961-1 (MUESTRAS)')
        response = self.client.get(reverse('universe'), {'catalog': 'all'})
        self.assertNotContains(response, duplicate.number)
        self.assertContains(response, '<b>3P848961-1</b>')
        self.assertEqual(response.context['catalog_total'], 2)
        self.assertTrue(Part.objects.filter(pk=duplicate.pk).exists())

    def test_corrected_carrier_alias_preserves_production_and_uses_normal_quantity(self):
        original = Part.objects.create(number='48TM004286B', client=self.customer, in_universe_ramos=True)
        duplicate = Part.objects.create(number='48TM004286', client=self.customer, in_universe_ramos=True)
        order = self.create(part_number=duplicate.number, quantity=14000)
        machine = Machine.objects.create(code='CARRIER-MACHINE')
        work = start_production(order=order, machine=machine, quantity=6, employee=self.user)
        close_production(work_item=work, quantity=3, employee=self.user)
        models = (ProductionOrder, WorkInProcess, ProductionClose, Movement, Inventory, InventoryBucket)
        before = {model: list(model.objects.order_by('pk').values()) for model in models}
        before_part = Part.objects.values().get(pk=duplicate.pk)
        link = import_module('operations.migrations.0020_link_corrected_carrier_part').link_corrected_carrier_part
        link(apps, SimpleNamespace(connection=connection))
        link(apps, SimpleNamespace(connection=connection))
        self.assertEqual(Part.objects.values().get(pk=duplicate.pk), {**before_part, 'sample_base_id': original.pk})
        for model in models:
            self.assertEqual(list(model.objects.order_by('pk').values()), before[model])
        order.refresh_from_db()
        self.assertEqual(order.part_display, original.number)
        self.assertFalse(order.is_sample)
        for number in (duplicate.number, original.number):
            new = self.create(part_number=number, quantity=37)
            self.assertEqual(new.part_id, original.pk)
            self.assertEqual(new.quantity, 37)
            self.assertFalse(new.is_sample)
        sample = self.create(part_number=duplicate.number, is_sample=True, quantity=37)
        self.assertEqual(sample.part_id, original.pk)
        self.assertEqual(sample.quantity, 5)
        self.assertTrue(sample.is_sample)
        response = self.client.get(reverse('universe'), {'catalog': 'all'})
        self.assertNotContains(response, '<b>48TM004286</b>')
        self.assertContains(response, '<b>48TM004286B</b>')

    def test_corrected_alias_in_bulk_preview_is_not_a_sample(self):
        original = Part.objects.create(number='48TM004286B', client=self.customer)
        Part.objects.create(number='48TM004286', client=self.customer, sample_base=original)
        wb = Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        ws.append(['ID Cliente', 'Orden de Produccion', 'Linea Prod Clte', 'Fecha de Entrega', 'Num. Parte', 'Cantidad', 'Linea'])
        ws.append(['S01', 'BULK-CORRECTED', '', None, '48TM004286', 14000, ''])
        stream = BytesIO()
        wb.save(stream)
        response = self.client.post(reverse('bulk-load-program'), {
            'file': SimpleUploadedFile('corrected.xlsx', stream.getvalue())})
        self.assertContains(response, original.number)
        token = response.context['preview_token']
        self.assertTrue(token)
        result = self.client.post(reverse('bulk-load-program'), {'action': 'confirm', 'preview_token': token})
        self.assertEqual(result.status_code, 302)
        order = ProductionOrder.objects.get(program='BULK-CORRECTED')
        self.assertEqual(order.part_id, original.pk)
        self.assertFalse(order.is_sample)
        self.assertEqual(order.quantity, 14000)
