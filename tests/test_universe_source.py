from io import BytesIO

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from openpyxl import Workbook, load_workbook

from operations.forms import UniversePartForm
from operations.models import ModuleAccess, Part
from operations.views import _import_universe_workbook


class UniverseSourceTests(TestCase):
    def source_file(self):
        workbook = Workbook()
        other = workbook.active
        other.title = "Sheet1"
        other.append(["CLIENTE", "NÚMERO DE PARTE", "DIAMETRO"])
        other.append(["Incorrecto", "WRONG", "1"])
        sheet = workbook.create_sheet("Sheet1 (2)")
        sheet.append([])
        sheet.append([None, "CLIENTE", "ID", "DESTINO", "NÚMERO DE PARTE",
                      "DIAMETRO", "PARED", "PESO", "CAJA", "STD.",
                      "HOJA COLOR", None, "CLIENTE", "PLANTA"])
        sheet.append([1, "LENNOX 1", "CD0036", "SAIRA MORALES", "625994-01",
                      "A - 3/8", 0.42, 0.02, "CTD T.MÉXICO", 1000,
                      "BLANCO", None, "Cliente secundario", 0])
        sheet.append([2, "RHEEM", "#N/A", "EZO", "82-104894-08",
                      "0.5", 0.028, 311])
        sheet.append([3, "RHEEM", "#N/A", "EZO", "82-104894-08",
                      "1/2", 0.028, 311])
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return output

    def test_import_uses_source_sheet_and_first_client_column(self):
        created, updated, skipped = _import_universe_workbook(self.source_file())

        self.assertEqual((created, updated, skipped), (2, 0, 0))
        self.assertFalse(Part.objects.filter(number="WRONG").exists())
        part = Part.objects.get(number="625994-01")
        self.assertTrue(part.in_universe_ramos)
        self.assertEqual(part.client.name, "LENNOX 1")
        self.assertEqual(part.diameter, "A - 3/8")
        rheem = Part.objects.get(number="82-104894-08")
        self.assertEqual(rheem.diameter, "1/2")
        self.assertEqual(rheem.client.external_id, "")

    def test_import_replaces_catalog_membership_without_deleting_parts(self):
        old = Part.objects.create(number="OLD-CATALOG", in_universe_ramos=True)
        operational = Part.objects.create(number="OPERATIONAL")

        _import_universe_workbook(self.source_file())

        old.refresh_from_db()
        operational.refresh_from_db()
        self.assertFalse(old.in_universe_ramos)
        self.assertFalse(operational.in_universe_ramos)
        self.assertEqual(Part.objects.filter(in_universe_ramos=True).count(), 2)
        self.assertEqual(Part.objects.count(), 4)

    def test_conflicting_diameters_are_not_guessed(self):
        workbook = load_workbook(self.source_file())
        sheet = workbook["Sheet1 (2)"]
        sheet.append([4, "RHEEM", "CD0082", "EZO", "P-CONFLICT", "1/2"])
        sheet.append([5, "RHEEM", "CD0082", "EZO", "P-CONFLICT", "3/4"])
        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        _import_universe_workbook(output)

        self.assertEqual(Part.objects.get(number="P-CONFLICT").diameter, "")

    def test_universe_does_not_show_description(self):
        self.assertNotIn("description", UniversePartForm().fields)
        user = get_user_model().objects.create_user("catalog-reader", password="secret-pass")
        access = ModuleAccess.objects.get(user=user)
        access.universe = True
        access.save(update_fields=["universe"])
        self.client.force_login(user)
        Part.objects.create(number="VISIBLE", in_universe_ramos=True)
        Part.objects.create(number="HIDDEN")
        response = self.client.get(reverse("universe"))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "VISIBLE")
        self.assertNotContains(response, "HIDDEN")
        self.assertNotContains(response, "Descripción")

    def test_can_switch_between_ramos_and_every_registered_part(self):
        user = get_user_model().objects.create_user("all-parts-reader", password="secret-pass")
        access = ModuleAccess.objects.get(user=user)
        access.universe = True
        access.save(update_fields=["universe"])
        self.client.force_login(user)
        Part.objects.create(number="RAMOS-PART", in_universe_ramos=True)
        Part.objects.bulk_create([Part(number=f"OTHER-{index:04d}") for index in range(1001)])

        ramos = self.client.get(reverse("universe"))
        self.assertEqual(ramos.context["universe_total"], 1)
        self.assertEqual(len(ramos.context["parts"]), 1)
        self.assertContains(ramos, "Ver todos los números (1002)")

        complete = self.client.get(reverse("universe"), {"catalog": "all"})
        self.assertEqual(complete.context["universe_total"], 1002)
        self.assertEqual(len(complete.context["parts"]), 1002)
        self.assertContains(complete, "OTHER-1000")
        self.assertContains(complete, "Ver solo Universo Ramos (1)")
        self.assertContains(complete, 'name="catalog" value="all"', count=2)

        filtered = self.client.get(reverse("universe"), {"catalog": "all", "q": "OTHER-1000"})
        self.assertEqual(len(filtered.context["parts"]), 1)
        self.assertEqual(filtered.context["parts"][0].number, "OTHER-1000")
        self.assertFalse(Part.objects.get(number="OTHER-1000").in_universe_ramos)
