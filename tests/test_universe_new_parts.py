from io import BytesIO
from datetime import timedelta

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from operations.models import ModuleAccess, Part


class NewUniversePartsTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user("parts-uploader", password="secret-pass")
        access = ModuleAccess.objects.get(user=self.user)
        access.universe = True
        access.universe_import = True
        access.universe_edit = True
        access.save(update_fields=["universe", "universe_import", "universe_edit"])
        self.client.force_login(self.user)

    def upload(self, rows):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1"
        sheet.append(["Número de parte", "Cliente", "Diámetro", "Peso unitario (kg)"])
        for row in rows:
            sheet.append(row)
        output = BytesIO()
        workbook.save(output)
        return SimpleUploadedFile(
            "piezas_nuevas.xlsx", output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    def test_download_template_and_append_new_parts_with_date_filter(self):
        template = self.client.get(reverse("download-universe-new-parts-template"))
        self.assertEqual(template.status_code, 200)
        workbook = load_workbook(BytesIO(template.content))
        self.assertEqual(workbook.sheetnames, ["Sheet1", "Instrucciones"])
        self.assertEqual(
            [workbook["Sheet1"].cell(1, column).value for column in range(1, 5)],
            ["Número de parte", "Cliente", "Diámetro", "Peso unitario (kg)"],
        )
        workbook.close()

        response = self.client.post(reverse("universe"), {
            "action": "import_new",
            "file": self.upload([
                ("NEW-001", "RHEEM", "1/2", 0.25),
                ("NEW-002", "RHEEM", "3/8", None),
            ]),
        })
        self.assertEqual(response.status_code, 302)
        self.assertEqual(Part.objects.filter(in_universe_ramos=True).count(), 2)
        part = Part.objects.get(number="NEW-001")
        self.assertIsNotNone(part.universe_added_at)
        self.assertEqual(str(part.unit_weight_kg), "0.250000")
        day = timezone.localtime(part.universe_added_at).date().isoformat()
        filtered = self.client.get(reverse("universe"), {"uploaded_on": day})
        self.assertEqual(len(filtered.context["parts"]), 2)
        self.assertContains(filtered, 'name="uploaded_on" value="' + day + '"')

    def test_duplicate_in_file_rejects_entire_upload(self):
        response = self.client.post(reverse("universe"), {
            "action": "import_new",
            "file": self.upload([
                ("SAME-1", "RHEEM", "1/2", None),
                ("same-1", "RHEEM", "3/8", None),
            ]),
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "repetido en el archivo")
        self.assertFalse(Part.objects.exists())

    def test_existing_number_outside_universe_rejects_entire_upload(self):
        Part.objects.create(number="OLD-1")
        response = self.client.post(reverse("universe"), {
            "action": "import_new",
            "file": self.upload([
                ("NEW-1", "RHEEM", "1/2", None),
                ("old-1", "RHEEM", "3/8", None),
            ]),
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ya está registrado")
        self.assertEqual(Part.objects.count(), 1)

    def test_invalid_row_rejects_entire_upload(self):
        response = self.client.post(reverse("universe"), {
            "action": "import_new",
            "file": self.upload([
                ("NEW-1", "RHEEM", "1/2", None),
                ("NEW-2", "RHEEM", "3/8", -1),
            ]),
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "peso unitario inválido")
        self.assertFalse(Part.objects.exists())

    def test_manual_add_rejects_case_insensitive_duplicate(self):
        Part.objects.create(number="ABC-123", in_universe_ramos=True)
        response = self.client.post(reverse("universe"), {
            "action": "add", "number": "abc-123", "client": "",
            "diameter": "", "unit_weight_kg": "",
        })
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "ya está registrado")
        self.assertEqual(Part.objects.count(), 1)

    def test_import_permission_is_required(self):
        access = ModuleAccess.objects.get(user=self.user)
        access.universe_import = False
        access.save(update_fields=["universe_import"])
        response = self.client.post(reverse("universe"), {
            "action": "import_new",
            "file": self.upload([("NEW-1", "RHEEM", "1/2", None)]),
        })
        self.assertEqual(response.status_code, 302)
        self.assertFalse(Part.objects.exists())

    def test_list_is_ordered_from_first_upload_to_last(self):
        now = timezone.now()
        Part.objects.create(
            number="Z-NEWEST", in_universe_ramos=True,
            universe_added_at=now,
        )
        Part.objects.create(
            number="A-OLDEST", in_universe_ramos=True,
            universe_added_at=now - timedelta(days=2),
        )
        Part.objects.create(
            number="M-MIDDLE", in_universe_ramos=True,
            universe_added_at=now - timedelta(days=1),
        )
        Part.objects.create(number="NO-UPLOAD-DATE")

        ramos = self.client.get(reverse("universe"))
        self.assertEqual(
            [part.number for part in ramos.context["parts"]],
            ["A-OLDEST", "M-MIDDLE", "Z-NEWEST"],
        )
        complete = self.client.get(reverse("universe"), {"catalog": "all"})
        self.assertEqual(
            [part.number for part in complete.context["parts"]],
            ["A-OLDEST", "M-MIDDLE", "Z-NEWEST", "NO-UPLOAD-DATE"],
        )
