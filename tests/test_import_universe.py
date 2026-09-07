from pathlib import Path
from tempfile import TemporaryDirectory

from django.core.management import call_command
from django.test import TestCase
from openpyxl import Workbook

from operations.models import Part


class ImportUniverseDiameterTests(TestCase):
    def make_workbook(self, rows):
        temp_dir = TemporaryDirectory()
        path = Path(temp_dir.name) / "universo.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sheet1 (2)"
        headers = [
            None, "CLIENTE", "ID", "DESTINO", "NÚMERO DE PARTE", "DIAMETRO",
            "PARED", "PESO", "CAJA", "STD.", "HOJA COLOR", None, "CLIENTE", "PLANTA",
        ]
        for column, value in enumerate(headers, 1):
            sheet.cell(row=2, column=column, value=value)
        for index, row in enumerate(rows, 1):
            client, external_id, part_number, diameter = row
            values = [
                index, client, external_id, "DESTINO", part_number, diameter,
                None, None, None, None, None, None, client, None,
            ]
            for column, value in enumerate(values, 1):
                sheet.cell(row=index + 2, column=column, value=value)
        workbook.save(path)
        return temp_dir, path

    def test_import_keeps_exact_text_and_only_fills_empty_diameters(self):
        empty = Part.objects.create(number="P-EMPTY")
        existing = Part.objects.create(number="P-EXISTING", diameter="7/8")
        conflict = Part.objects.create(number="P-CONFLICT")
        temp_dir, path = self.make_workbook([
            ("Cliente Uno", "C1", empty.number, "A - 3/8"),
            ("Cliente Uno", "C1", existing.number, "0.875"),
            ("Cliente Uno", "C1", conflict.number, "0.5"),
            ("Cliente Uno", "C1", conflict.number, "1/2"),
        ])
        self.addCleanup(temp_dir.cleanup)

        call_command("import_universe", path)

        empty.refresh_from_db()
        existing.refresh_from_db()
        conflict.refresh_from_db()
        self.assertEqual(empty.diameter, "A - 3/8")
        self.assertEqual(existing.diameter, "7/8")
        self.assertEqual(conflict.diameter, "")

    def test_overwrite_requires_explicit_option(self):
        part = Part.objects.create(number="P-EDITABLE", diameter="Original manual")
        temp_dir, path = self.make_workbook([
            ("Cliente Uno", "C1", part.number, "Texto del Excel"),
        ])
        self.addCleanup(temp_dir.cleanup)

        call_command("import_universe", path, overwrite_diameters=True)

        part.refresh_from_db()
        self.assertEqual(part.diameter, "Texto del Excel")

    def test_dry_run_reports_but_does_not_save(self):
        part = Part.objects.create(number="P-DRY-RUN")
        temp_dir, path = self.make_workbook([
            ("Cliente Uno", "C1", part.number, "C - 1/2"),
        ])
        self.addCleanup(temp_dir.cleanup)

        call_command("import_universe", path, dry_run=True)

        part.refresh_from_db()
        self.assertEqual(part.diameter, "")
        self.assertIsNone(part.client)

    def test_diameters_only_does_not_modify_client_relationship(self):
        part = Part.objects.create(number="P-DIAMETER-ONLY")
        temp_dir, path = self.make_workbook([
            ("Cliente Uno", "C1", part.number, "5/8"),
        ])
        self.addCleanup(temp_dir.cleanup)

        call_command("import_universe", path, diameters_only=True)

        part.refresh_from_db()
        self.assertEqual(part.diameter, "5/8")
        self.assertIsNone(part.client)
