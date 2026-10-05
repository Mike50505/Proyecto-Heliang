"""Append new parts to Universo Ramos from its dedicated Excel template."""

import re
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.core.validators import DecimalValidator
from django.db import transaction
from django.db.models.functions import Lower
from django.utils import timezone
from openpyxl import load_workbook

from .models import Client, Part
from .security import MAX_XLSX_ROWS, validate_xlsx_archive


HEADERS = ("Número de parte", "Cliente", "Diámetro", "Peso unitario (kg)")


def cell_text(value):
    return "" if value is None else str(value).strip()


@transaction.atomic
def import_new_universe_parts(upload):
    validate_xlsx_archive(upload)
    workbook = load_workbook(upload, data_only=True, read_only=True)
    try:
        if "Sheet1" not in workbook.sheetnames:
            raise ValidationError("La plantilla debe contener una hoja llamada Sheet1.")
        sheet = workbook["Sheet1"]
        headers = tuple(cell_text(sheet.cell(1, column).value) for column in range(1, 5))
        if headers != HEADERS:
            raise ValidationError("Los encabezados deben ser: " + ", ".join(HEADERS) + ".")

        rows = []
        seen = set()
        errors = []
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            if row_number > MAX_XLSX_ROWS + 1:
                raise ValidationError(f"El archivo excede el máximo de {MAX_XLSX_ROWS} filas.")
            number, client_name, diameter, raw_weight = (
                values[index] if index < len(values) else None for index in range(4)
            )
            if all(value is None or str(value).strip() == "" for value in
                   (number, client_name, diameter, raw_weight)):
                continue
            number = cell_text(number)
            client_name = cell_text(client_name)
            diameter = cell_text(diameter)
            if not number or not client_name:
                errors.append(f"Fila {row_number}: número de parte y cliente son obligatorios.")
                continue
            if len(number) > 80 or len(client_name) > 120 or len(diameter) > 40:
                errors.append(f"Fila {row_number}: un texto supera el largo permitido.")
                continue
            key = number.lower()
            if key in seen:
                errors.append(f"Fila {row_number}: el número {number} está repetido en el archivo.")
                continue
            seen.add(key)
            weight = None
            if raw_weight is not None and str(raw_weight).strip():
                try:
                    weight = Decimal(str(raw_weight).replace(",", ""))
                    DecimalValidator(14, 6)(weight)
                    if not weight.is_finite() or weight < 0:
                        raise ValueError
                except (InvalidOperation, ValueError, ValidationError):
                    errors.append(f"Fila {row_number}: peso unitario inválido.")
                    continue
            rows.append((row_number, number, client_name, diameter, weight))
        if not rows and not errors:
            errors.append("La plantilla no contiene piezas para importar.")
        existing = set(
            Part.objects.annotate(number_lower=Lower("number"))
            .filter(number_lower__in=seen).values_list("number_lower", flat=True)
        )
        for row_number, number, *_ in rows:
            if number.lower() in existing:
                errors.append(f"Fila {row_number}: el número {number} ya está registrado.")
        if errors:
            raise ValidationError(errors)

        added_at = timezone.now()
        for _, number, client_name, diameter, weight in rows:
            code = re.sub(r"[^A-Z0-9]+", "-", client_name.upper()).strip("-")[:30] or "SIN-CLIENTE"
            client = Client.objects.filter(name__iexact=client_name).first()
            if client is None:
                client = Client.objects.filter(code=code).first()
                if client and client.name.casefold() != client_name.casefold():
                    raise ValidationError(
                        f"El cliente {client_name} coincide con el código de {client.name}.")
                if client is None:
                    client = Client.objects.create(code=code, name=client_name)
            Part.objects.create(
                number=number, client=client, diameter=diameter,
                unit_weight_kg=weight, in_universe_ramos=True,
                universe_added_at=added_at,
            )
        return len(rows)
    finally:
        workbook.close()
